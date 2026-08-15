"""Reads a source document's own metadata — the §5.1 credibility signals
`author` and `last_modified` — at parse time.

**Why this isn't Docling.** It lives beside the Docling parser because it is
part of the same port's job, but Docling exposes none of this: its
`DoclingDocument.origin` carries only mimetype, binary hash, filename and uri,
and nothing on the conversion result or the PDF backend surfaces an author or a
date. So each format is read directly, with the libraries Docling already pulls
in (pypdfium2, python-docx, python-pptx, openpyxl).

**Why there is no filesystem-mtime fallback.** A file's mtime is when it landed
in the inbox, not when the source last changed — using it would stamp every
hand-dropped note as freshly updated and make the recency signal a lie. ADR
0001 is explicit that absent means *unknown* and unknown is neutral, which is
strictly better than a fabricated date. Signals we cannot read stay `None`.
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from pathlib import Path

from pipeline.domain.source_document import DocumentMetadata

logger = logging.getLogger(__name__)

# PDF date syntax (PDF 32000-1 §7.9.4): D:YYYYMMDDHHmmSSOHH'mm', where every
# part after the year is optional and O is one of +-Z.
_PDF_DATE_PATTERN = re.compile(r"^D:(\d{4})(\d{2})?(\d{2})?")


def read_document_metadata(path: str) -> DocumentMetadata:
    """Best-effort. A document that carries no metadata, or a reader that
    fails on it, yields empty signals rather than an error — a missing
    credibility signal must never be a reason a document fails to ingest."""
    suffix = Path(path).suffix.lower()
    try:
        if suffix == ".pdf":
            return _read_pdf(path)
        if suffix == ".docx":
            return _read_docx(path)
        if suffix == ".pptx":
            return _read_pptx(path)
        if suffix == ".xlsx":
            return _read_xlsx(path)
    except Exception:  # noqa: BLE001 - see docstring: signals are never fatal
        logger.warning("could not read source metadata from %s", path, exc_info=True)
    return DocumentMetadata()


def _read_pdf(path: str) -> DocumentMetadata:
    import pypdfium2 as pdfium

    info = pdfium.PdfDocument(path).get_metadata_dict()
    # ModDate is the real recency signal; CreationDate stands in for a document
    # that was never revised, which is still the document's own claim about
    # itself rather than something about our copy of it.
    raw_date = _clean(info.get("ModDate")) or _clean(info.get("CreationDate"))
    return DocumentMetadata(
        author=_clean(info.get("Author")),
        last_modified=_parse_pdf_date(raw_date),
    )


def _read_docx(path: str) -> DocumentMetadata:
    import docx

    return _from_core_properties(docx.Document(path).core_properties)


def _read_pptx(path: str) -> DocumentMetadata:
    import pptx

    return _from_core_properties(pptx.Presentation(path).core_properties)


def _read_xlsx(path: str) -> DocumentMetadata:
    import openpyxl

    properties = openpyxl.load_workbook(path, read_only=True).properties
    return DocumentMetadata(
        author=_clean(properties.creator),
        last_modified=_as_date(properties.modified) or _as_date(properties.created),
    )


def _from_core_properties(properties: object) -> DocumentMetadata:
    """OOXML core properties, shared shape across python-docx and python-pptx."""
    return DocumentMetadata(
        author=_clean(getattr(properties, "author", None)),
        last_modified=(
            _as_date(getattr(properties, "modified", None))
            or _as_date(getattr(properties, "created", None))
        ),
    )


def _clean(value: object) -> str | None:
    """Empty strings are how every one of these formats spells "not set" — a
    blank author must read as unknown, not as an author named ""."""
    if not isinstance(value, str):
        return None
    stripped = value.strip()
    return stripped or None


def _as_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def _parse_pdf_date(value: str | None) -> str | None:
    """`D:20240410211143Z` -> `2024-04-10`. §5.1 specifies `YYYY-MM-DD`, so the
    time and offset are dropped rather than carried."""
    if value is None:
        return None
    match = _PDF_DATE_PATTERN.match(value)
    if match is None:
        return None
    year, month, day = match.group(1), match.group(2) or "01", match.group(3) or "01"
    try:
        return date(int(year), int(month), int(day)).isoformat()
    except ValueError:
        return None
